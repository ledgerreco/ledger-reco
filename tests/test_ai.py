"""AI layer tests with a fake Grok client (no network)."""
from pathlib import Path
import io

import openpyxl
import pandas as pd

import ai
import reco

S = Path(__file__).resolve().parent.parent / "samples"


def _load(name):
    return reco.ledger_from_file((S / name).read_bytes(), name)[0]


def _rule_v2():
    return reco.reconcile(_load("ai_our_books_tally_v2.xlsx"), _load("ai_party_statement_v2.xlsx"),
                          label_a="Our Books", label_b="Sharma")


def _id(res, side, ref):
    df = res.A if side == "A" else res.B
    return f"{side}{df.index[df['ref'] == ref][0]}"


class FakeClient:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []
        self.usage = ai.Usage()

    def call_json(self, system, content, schema, name, effort=None):
        self.calls.append((name, content))
        r = self.responses.pop(0)
        if isinstance(r, Exception):
            raise r
        return r


def test_rule_v2_baseline():
    s = _rule_v2().summary
    assert (s["matched_count"], s["only_a_count"], s["only_b_count"], s["mismatch_count"]) == (13, 4, 4, 2)


def test_ai_groups_validated_and_applied():
    rule = _rule_v2()
    g_pay = {"a_ids": [_id(rule, "A", "34")], "b_ids": [_id(rule, "B", "RCPT-0091"), _id(rule, "B", "RCPT-0092")],
             "kind": "one_to_many", "explained_difference": 0, "difference_reason": "", "reason": "one NEFT for 2 bills",
             "confidence": 0.95}
    g_split = {"a_ids": [_id(rule, "A", "SHT/113"), _id(rule, "A", "SHT/113A")], "b_ids": [_id(rule, "B", "SI-113")],
               "kind": "split_entry", "explained_difference": 0, "difference_reason": "", "reason": "invoice split",
               "confidence": 0.9}
    bad_sum = {"a_ids": [_id(rule, "A", "SHT/108")], "b_ids": [_id(rule, "B", "SI-110")], "kind": "other",
               "explained_difference": 0, "difference_reason": "", "reason": "guess", "confidence": 0.9}
    bad_id = {"a_ids": ["A999"], "b_ids": [_id(rule, "B", "INT/09")], "kind": "other", "explained_difference": 0,
              "difference_reason": "", "reason": "x", "confidence": 0.9}
    reuse = {**g_pay, "confidence": 0.7}
    fake = FakeClient([
        {"groups": [g_pay, g_split, bad_sum, bad_id, reuse],
         "notes": [{"ids": [_id(rule, "A", "SHT/108")], "category": "possible_duplicate", "note": "108 booked twice",
                    "suggested_action": "Reverse duplicate"}]},
        {"summary_en": "ok", "summary_hi": "ठीक", "action_items": [{"action": "Ask for DN/04", "priority": "high"}]},
    ])
    res = ai.run_ai(rule, fake)
    s = res.summary
    assert res.ai_status == "ok"
    assert s["ai_group_count"] == 2
    assert (s["only_a_count"], s["only_b_count"], s["mismatch_count"]) == (2, 2, 1)  # 108-dup, DN/04 | SI-110, INT/09
    assert abs(res.statement.iloc[-1]["Amount"]) < 0.01
    cats = res.ai_notes["Category"].tolist()
    assert cats.count("ai_proposal_rejected") == 3 and "possible_duplicate" in cats
    assert res.ai_summary["action_items"][0]["priority"] == "high"
    wb = openpyxl.load_workbook(io.BytesIO(reco.to_excel(res)))
    assert {"AI Matched", "AI Notes"} <= set(wb.sheetnames)
    text = " ".join(str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value)
    assert "AI summary" in text and "Ask for DN/04" in text


def test_explained_difference_tds():
    rule = _rule_v2()
    # SHT/103 23,600 vs SI-103 24,600 : difference 1,000 – pretend "GST difference" (4 %, within 10 %)
    g = {"a_ids": [_id(rule, "A", "SHT/103")], "b_ids": [_id(rule, "B", "SI-103")], "kind": "tds_gst_difference",
         "explained_difference": 1000.0, "difference_reason": "GST charged on freight", "reason": "same invoice",
         "confidence": 0.8}
    acc, rej = ai.validate_groups(rule, [g], set(rule.A.index), set(rule.B.index), 1.0)
    assert len(acc) == 1 and acc[0]["status"].startswith("Matched – difference")
    g2 = {**g, "explained_difference": 0.0}
    acc, rej = ai.validate_groups(rule, [g2], set(rule.A.index), set(rule.B.index), 1.0)
    assert not acc and "do not agree" in rej[0]["rejected_because"]
    g3 = {**g, "kind": "other"}  # "amounts differ" without a TDS/GST/discount/rounding cause is not a match
    acc, rej = ai.validate_groups(rule, [g3], set(rule.A.index), set(rule.B.index), 1.0)
    assert not acc


def test_ai_failure_keeps_rule_result():
    rule = _rule_v2()
    res = ai.run_ai(rule, FakeClient([TimeoutError("slow")]))
    assert res.ai_status.startswith("error") and res.summary["matched_count"] == 13
    reco.to_excel(res)  # still exports


def test_chunking():
    rule = _rule_v2()
    ids_a = sorted(rule.free_a | {i for i, _ in rule.mism_pairs})
    ids_b = sorted(rule.free_b | {j for _, j in rule.mism_pairs})
    chunks = list(ai._chunks(rule, ids_a, ids_b, max_rows=4))
    assert sum(len(a) + len(b) for a, b in chunks) == len(ids_a) + len(ids_b)
    assert all(len(a) + len(b) <= 4 for a, b in chunks)


def test_vision_grid_and_render():
    imgs = ai.render_pages((S / "scanned_party_statement.pdf").read_bytes(), "x.pdf")
    assert len(imgs) == 1 and imgs[0][:4] == b"\x89PNG"
    fake = FakeClient([{"readable": True, "opening_balance": 50000, "opening_balance_side": "Dr", "rows": [
        {"date": "2024-04-05", "ref": "SI-101", "narration": "Sales Invoice", "debit": 118000, "credit": None},
        {"date": "2024-04-22", "ref": "RCPT-0045", "narration": "Receipt", "debit": None, "credit": 100000}]}])
    grid, info = ai.ai_extract_rows(imgs, fake)
    df = reco.apply_header(grid, reco.detect_header_row(grid))
    led = reco.standardize(df, reco.detect_columns(df))
    assert len(led.txns) == 2 and led.opening == 50000 and led.txns["signed"].sum() == 18000
    assert fake.calls[0][1][1]["type"] == "input_image"


def test_image_file_needs_vision():
    import pytest
    with pytest.raises(reco.NeedsVisionError):
        reco.load_file((S / "scanned_party_statement.jpg").read_bytes(), "x.jpg")


def test_api_key_lookup(monkeypatch):
    monkeypatch.delenv("XAI_API_KEY", raising=False)
    assert ai.get_api_key({}) is None
    assert ai.get_api_key({"XAI_API_KEY": "k1"}) == "k1"
    monkeypatch.setenv("XAI_API_KEY", "k2")
    assert ai.get_api_key(None) == "k2"


def test_instruction_passed_to_ai_calls():
    rule = _rule_v2()
    fake = FakeClient([{"groups": [], "notes": []},
                       {"summary_en": "s", "summary_hi": "स", "action_items": []}])
    systems = []
    orig = fake.call_json
    fake.call_json = lambda system, *a, **k: (systems.append(system), orig(system, *a, **k))[1]
    ai.run_ai(rule, fake, instruction="TDS 2% ignore karo")
    assert len(systems) == 2 and all("TDS 2% ignore karo" in s for s in systems)


def test_chat_context_and_reply():
    rule = _rule_v2()
    ctx = ai.chat_context(rule, "sirf March")
    assert "Reconciliation statement" in ctx and "all 19 entries" in ctx and "sirf March" in ctx
    assert len(ctx) <= 60000
    assert "No reconciliation" in ai.chat_context(None)

    class ChatFake:
        def call_text(self, messages, effort=None, max_output_tokens=0):
            self.messages = messages
            return "जवाब"
    cf = ChatFake()
    hist = [{"role": "user", "content": f"q{i}"} if i % 2 == 0 else {"role": "assistant", "content": "a" * 5000}
            for i in range(30)]
    assert ai.chat_reply(cf, ctx, hist, "SI-110 ka kya karna hai?") == "जवाब"
    assert cf.messages[0]["role"] == "system" and cf.messages[-1]["content"].startswith("SI-110")
    assert len(cf.messages) == 12 and all(len(m["content"]) <= 2500 for m in cf.messages[1:])
